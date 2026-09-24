import contextlib
import csv
import io
import json
import logging
import sys
from pathlib import Path

from openpyxl import load_workbook

from roadmapper import cli, schema
from tests.conftest import FixtureSource


def test_extract_writes_all_outputs(tmp_path, fixture_source, monkeypatch):
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    rc = cli.main(
        ["extract", "--lat", "-32.0018629", "--lon", "115.8924599", "--radius-km", "1.0", "--out", str(out)]
    )
    assert rc == 0
    for name in ("roads.csv", "roads_vertices.csv", "metadata.json", "roads.xlsx", "qa_plot.png"):
        assert (out / name).exists(), name
    with (out / "roads.csv").open(newline="") as fh:
        header = next(csv.reader(fh))
    assert header == schema.OUTPUT_COLUMNS
    meta = json.loads((out / "metadata.json").read_text())
    assert meta["radius_km"] == 1.0 and meta["segments_intersecting"] > 0
    assert load_workbook(out / "roads.xlsx", read_only=True).sheetnames[0] == "roads"


def test_no_plot_flag(tmp_path, fixture_source, monkeypatch):
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    rc = cli.main(
        [
            "extract",
            "--lat",
            "-32.0018629",
            "--lon",
            "115.8924599",
            "--radius-km",
            "1.0",
            "--out",
            str(out),
            "--no-plot",
        ]
    )
    assert rc == 0
    assert not (out / "qa_plot.png").exists()
    assert (out / "roads.csv").exists()


def test_no_plot_flag_removes_stale_png_from_a_previous_run(tmp_path, fixture_source, monkeypatch, caplog):
    """--no-plot after a prior run into the same --out must not leave that run's PNG behind:
    it would silently disagree with the new CSVs sitting next to it while looking current."""
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    out.mkdir(parents=True)
    (out / "qa_plot.png").write_bytes(b"stale")
    with caplog.at_level("INFO"):
        rc = cli.main(
            [
                "extract",
                "--lat",
                "-32.0018629",
                "--lon",
                "115.8924599",
                "--radius-km",
                "1.0",
                "--out",
                str(out),
                "--no-plot",
            ]
        )
    assert rc == 0
    assert not (out / "qa_plot.png").exists()
    assert (out / "roads.csv").exists()
    assert any(
        "removed stale" in r.message and "qa_plot.png" in r.message for r in caplog.records
    )


def test_no_xlsx_flag_removes_stale_xlsx_from_a_previous_run(tmp_path, fixture_source, monkeypatch, caplog):
    """--no-xlsx after a prior run into the same --out must not leave that run's workbook behind."""
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    out.mkdir(parents=True)
    (out / "roads.xlsx").write_bytes(b"stale")
    with caplog.at_level("INFO"):
        rc = cli.main(
            [
                "extract",
                "--lat",
                "-32.0018629",
                "--lon",
                "115.8924599",
                "--radius-km",
                "1.0",
                "--out",
                str(out),
                "--no-xlsx",
            ]
        )
    assert rc == 0
    assert not (out / "roads.xlsx").exists()
    assert (out / "roads.csv").exists()
    assert any(
        "removed stale" in r.message and "roads.xlsx" in r.message for r in caplog.records
    )


def test_plot_failure_is_best_effort(tmp_path, fixture_source, monkeypatch, caplog):
    """A plotting exception must never fail the run: the CSVs are the deliverable."""
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)

    def _boom(*args, **kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr("roadmapper.plot.render_qa_plot", _boom)
    out = tmp_path / "run"
    # A stale PNG from an earlier run into this same --out (or a truncated one from a savefig
    # that died partway) must not survive a failed plot: it would no longer be a rendering of
    # what this run just wrote.
    out.mkdir(parents=True)
    (out / "qa_plot.png").write_bytes(b"stale")
    with caplog.at_level("WARNING"):
        rc = cli.main(
            [
                "extract",
                "--lat",
                "-32.0018629",
                "--lon",
                "115.8924599",
                "--radius-km",
                "1.0",
                "--out",
                str(out),
            ]
        )
    assert rc == 0
    assert (out / "roads.csv").exists()
    assert (out / "metadata.json").exists()
    assert not (out / "qa_plot.png").exists()
    assert any("QA plot not written" in r.message for r in caplog.records)


def test_plot_failure_when_matplotlib_absent(tmp_path, fixture_source, monkeypatch, caplog):
    """Simulate matplotlib/roadmapper.plot being unimportable; the run must still succeed."""
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    monkeypatch.setitem(sys.modules, "roadmapper.plot", None)
    out = tmp_path / "run"
    with caplog.at_level("WARNING"):
        rc = cli.main(
            [
                "extract",
                "--lat",
                "-32.0018629",
                "--lon",
                "115.8924599",
                "--radius-km",
                "1.0",
                "--out",
                str(out),
            ]
        )
    assert rc == 0
    assert (out / "roads.csv").exists()
    assert (out / "metadata.json").exists()
    assert not (out / "qa_plot.png").exists()
    assert any("QA plot not written" in r.message for r in caplog.records)


def test_no_xlsx_flag(tmp_path, fixture_source, monkeypatch):
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    assert (
        cli.main(
            [
                "extract",
                "--lat",
                "-32.0018629",
                "--lon",
                "115.8924599",
                "--radius-km",
                "1.0",
                "--out",
                str(out),
                "--no-xlsx",
            ]
        )
        == 0
    )
    assert not (out / "roads.xlsx").exists()
    # The user's example command is `extract --no-xlsx`: the plot must still be written, since
    # with --no-xlsx the CSVs (and the PNG) are the whole output.
    assert (out / "qa_plot.png").exists()


def test_to_xlsx_subcommand(tmp_path, fixture_source, monkeypatch):
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    cli.main(
        [
            "extract",
            "--lat",
            "-32.0018629",
            "--lon",
            "115.8924599",
            "--radius-km",
            "1.0",
            "--out",
            str(out),
            "--no-xlsx",
        ]
    )
    rc = cli.main(
        [
            "to-xlsx",
            "--roads",
            str(out / "roads.csv"),
            "--vertices",
            str(out / "roads_vertices.csv"),
            "--metadata",
            str(out / "metadata.json"),
            "--out",
            str(out / "converted.xlsx"),
        ]
    )
    assert rc == 0 and (out / "converted.xlsx").exists()


def test_no_roads_returns_exit_code_1(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: FixtureSource({17: [], 12: [], 16: [], 8: []}))
    rc = cli.main(
        ["extract", "--lat", "-32.0", "--lon", "115.89", "--radius-km", "1.0", "--out", str(tmp_path / "x")]
    )
    assert rc == 1
    assert "no road segments" in capsys.readouterr().err.lower()


def test_negative_radius_rejected(tmp_path):
    rc = cli.main(
        ["extract", "--lat", "-32.0", "--lon", "115.89", "--radius-km", "-1", "--out", str(tmp_path / "x")]
    )
    assert rc == 2


def test_osm_flag_fails_loudly_rather_than_silently_doing_nothing(
    tmp_path, fixture_source, monkeypatch, capsys
):
    """Passing --osm must be an error, not a no-op that reports success.

    Before this, the flag was advertised in the help, did nothing, and `metadata.json` still
    reported `osm_enabled: true` — so a user investigating the 5.55% width coverage could
    reasonably conclude OSM had been consulted and had nothing to add.
    """
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    rc = cli.main(["extract", "--lat", "-32.0018629", "--lon", "115.8924599", "--radius-km", "1.0",
                   "--out", str(out), "--osm"])
    assert rc == 1
    err = capsys.readouterr().err
    assert "not implemented" in err.lower()
    assert not (out / "roads.csv").exists(), "nothing may be written on a refused run"


def test_verbose_does_not_leak_matplotlib_debug_noise(tmp_path, fixture_source, monkeypatch, caplog):
    """-v must not be buried under matplotlib's font-scoring DEBUG spam.

    Before this fix, `extract ... -v` produced thousands of `matplotlib.font_manager` DEBUG
    lines (root logger at DEBUG propagates into matplotlib's loggers), drowning the roughly 66
    lines of the extraction's own logging. matplotlib's warnings/errors must still come through.

    Two pieces of process-global state would otherwise make this test pass even without the
    fix, if it ran after another test that already exercised the plot: matplotlib memoizes its
    font lookups for the life of the process (`fontManager._findfont_cached`, an `lru_cache`),
    so a warm cache silently skips the very debug calls this test checks for; and the
    `matplotlib` logger's level is itself global state that a previous *correctly fixed* call
    to `main()` would have already set to WARNING, which would mask a mutant that deletes the
    fix from THIS call. Both are reset below so the test's outcome depends only on what this
    call to `main()` does.
    """
    import matplotlib.font_manager

    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    mpl_logger = logging.getLogger("matplotlib")
    original_level = mpl_logger.level
    mpl_logger.setLevel(logging.NOTSET)
    matplotlib.font_manager.fontManager._findfont_cached.cache_clear()
    out = tmp_path / "run"
    try:
        with caplog.at_level(logging.DEBUG):
            rc = cli.main(
                [
                    "extract",
                    "--lat",
                    "-32.0018629",
                    "--lon",
                    "115.8924599",
                    "--radius-km",
                    "1.0",
                    "--out",
                    str(out),
                    "-v",
                ]
            )
        # Read the logger's OWN explicitly set level, and read it here, before the `finally`
        # restores it. It was set to NOTSET just above, so it can only now be WARNING if this
        # call to main() set it. An earlier version checked getEffectiveLevel() after the
        # restore instead, which inherited the root logger's restored WARNING and so could never
        # fail — the Opus re-review proved it passed with the fix removed. `.level` depends on
        # neither the root logger, caplog, nor matplotlib's font cache.
        level_set_by_main = mpl_logger.level
    finally:
        mpl_logger.setLevel(original_level)
    assert rc == 0
    assert level_set_by_main >= logging.WARNING, "main() must quieten matplotlib under -v"
    matplotlib_records = [r for r in caplog.records if r.name.startswith("matplotlib")]
    assert all(r.levelno >= logging.WARNING for r in matplotlib_records)
    roadmapper_records = [r for r in caplog.records if r.name.startswith("roadmapper")]
    assert any(r.levelno in (logging.DEBUG, logging.INFO) for r in roadmapper_records)


def test_osm_help_text_says_it_is_not_implemented():
    parser = cli.build_parser()
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.suppress(SystemExit):
        parser.parse_args(["extract", "--help"])
    assert "NOT IMPLEMENTED" in buf.getvalue()


def test_failure_to_remove_a_stale_file_warns_but_does_not_fail_the_run(
    tmp_path, fixture_source, monkeypatch, caplog
):
    """Tidying up is best-effort: a stale file that cannot be removed must not turn a successful
    extraction into an exit-1 failure after the CSVs are written — but it must be named loudly,
    since a stale file left in place is exactly the misleading artifact the removal prevents."""
    monkeypatch.setattr(cli, "MRWAClient", lambda **kwargs: fixture_source)
    out = tmp_path / "run"
    out.mkdir(parents=True)
    stale = out / "qa_plot.png"
    stale.write_bytes(b"stale")
    real_unlink = Path.unlink

    def refuse(self, *args, **kwargs):
        if self.name == "qa_plot.png":
            raise PermissionError("simulated: cannot remove")
        return real_unlink(self, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", refuse)
    with caplog.at_level("WARNING"):
        rc = cli.main(["extract", "--lat", "-32.0018629", "--lon", "115.8924599", "--radius-km", "1.0",
                       "--out", str(out), "--no-plot", "--no-xlsx"])
    assert rc == 0, "a housekeeping failure must not fail a successful extraction"
    assert (out / "roads.csv").exists()
    assert "does NOT describe this run" in caplog.text


def test_swapped_lat_lon_is_rejected_with_a_hint(tmp_path):
    """Passing (lon, lat) used to reach pyproj and die with a CRSError traceback."""
    err = io.StringIO()
    with contextlib.redirect_stderr(err):
        swapped = ["--lat", "115.8924599", "--lon", "-32.0018629"]
        rc = cli.main(["extract", *swapped, "--radius-km", "1", "--out", str(tmp_path / "r")])
    assert rc == 2
    assert "wrong order" in err.getvalue()
    assert not (tmp_path / "r").exists()


def test_non_finite_coordinates_and_radius_are_rejected(tmp_path):
    cases = (
        ["--lat", "nan", "--lon", "115.89", "--radius-km", "1"],
        ["--lat", "-32.0", "--lon", "inf", "--radius-km", "1"],
        ["--lat", "-32.0", "--lon", "115.89", "--radius-km", "inf"],
        ["--lat", "-32.0", "--lon", "200", "--radius-km", "1"],
        ["--lat", "95", "--lon", "115.89", "--radius-km", "1"],
    )
    for args in cases:
        err = io.StringIO()
        with contextlib.redirect_stderr(err):
            rc = cli.main(["extract", *args, "--out", str(tmp_path / "r")])
        assert rc == 2, args
        assert "usage:" in err.getvalue()
