import csv
import json

import pytest
from openpyxl import load_workbook

from roadmapper import schema
from roadmapper.export import EXCEL_MAX_CELL_CHARS, csv_to_xlsx, write_csv, write_json


def _road_row(i, wkt="LINESTRING (115.89 -32, 115.9 -32)"):
    row = dict.fromkeys(schema.OUTPUT_COLUMNS)
    row.update(
        ROAD="H001",
        ROAD_NAME="Albany Hwy",
        CWY="Left",
        START_SLK=0.0,
        END_SLK=0.3,
        OBJECTID=i,
        NETWORK_ELEMENT=f"H001/{i}-L",
        GlobalID="{ABC}",
        GEOMETRY_WKT=wkt,
        WIDTH_SOURCE="none",
        INSIDE_FRACTION=1.0,
        DATA_SOURCE=schema.DATA_SOURCE,
        EXTRACTED_AT_UTC="2026-09-22T00:00:00+00:00",
    )
    return row


def _vertex_rows(n):
    return [
        {
            "OBJECTID": 1,
            "NETWORK_ELEMENT": "H001/1-L",
            "ROAD": "H001",
            "ROAD_NAME": "Albany Hwy",
            "CWY": "Left",
            "PART": 0,
            "SEQ": s,
            "LAT": -32.0,
            "LON": 115.89 + s * 1e-4,
        }
        for s in range(n)
    ]


def test_write_csv_header_order_and_none_as_empty(tmp_path):
    path = write_csv([_road_row(1)], schema.OUTPUT_COLUMNS, tmp_path / "roads.csv")
    with path.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == schema.OUTPUT_COLUMNS
    assert rows[1][schema.OUTPUT_COLUMNS.index("WIDTH_M")] == ""
    assert rows[1][0] == "H001"


def test_write_csv_zero_rows_still_writes_header(tmp_path):
    path = write_csv([], schema.VERTEX_COLUMNS, tmp_path / "v.csv")
    assert path.read_text().strip() == ",".join(schema.VERTEX_COLUMNS)


def test_csv_to_xlsx_sheets_and_counts(tmp_path):
    roads = tmp_path / "roads.csv"
    verts = tmp_path / "roads_vertices.csv"
    meta = tmp_path / "metadata.json"
    write_csv([_road_row(1), _road_row(2)], schema.OUTPUT_COLUMNS, roads)
    write_csv(_vertex_rows(5), schema.VERTEX_COLUMNS, verts)
    write_json({"radius_km": 10, "network_type_counts": {"Local Road": 2}}, meta)
    report = csv_to_xlsx(roads, verts, meta, tmp_path / "roads.xlsx")
    wb = load_workbook(tmp_path / "roads.xlsx", read_only=True)
    assert wb.sheetnames == ["roads", "vertices", "metadata"]
    ws = wb["roads"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == schema.OUTPUT_COLUMNS
    assert ws.max_row == 3 and wb["vertices"].max_row == 6
    assert report == {
        "sheets": ["roads", "vertices", "metadata"],
        "roads_rows": 2,
        "vertices_rows": 5,
        "truncated_cells": 0,
    }
    meta_rows = {r[0].value: r[1].value for r in wb["metadata"].iter_rows(min_row=2)}
    assert meta_rows["radius_km"] in (10, "10")
    assert json.loads(meta_rows["network_type_counts"]) == {"Local Road": 2}


def test_csv_to_xlsx_keeps_ids_as_text(tmp_path):
    row = _road_row(1)
    row.update(ROAD="1010001", LG_NO="101")
    write_csv([row], schema.OUTPUT_COLUMNS, tmp_path / "r.csv")
    write_csv([], schema.VERTEX_COLUMNS, tmp_path / "v.csv")
    write_json({}, tmp_path / "m.json")
    csv_to_xlsx(tmp_path / "r.csv", tmp_path / "v.csv", tmp_path / "m.json", tmp_path / "o.xlsx")
    ws = load_workbook(tmp_path / "o.xlsx", read_only=True)["roads"]
    values = [c.value for c in list(ws.iter_rows(min_row=2, max_row=2))[0]]
    assert values[schema.OUTPUT_COLUMNS.index("ROAD")] == "1010001"
    assert values[schema.OUTPUT_COLUMNS.index("LG_NO")] == "101"


def test_csv_to_xlsx_splits_vertices_and_truncates_long_cells(tmp_path):
    long_wkt = "LINESTRING (" + ", ".join(f"115.{i:05d} -32.0" for i in range(4000)) + ")"
    assert len(long_wkt) > EXCEL_MAX_CELL_CHARS
    write_csv([_road_row(1, wkt=long_wkt)], schema.OUTPUT_COLUMNS, tmp_path / "r.csv")
    write_csv(_vertex_rows(7), schema.VERTEX_COLUMNS, tmp_path / "v.csv")
    write_json({}, tmp_path / "m.json")
    report = csv_to_xlsx(
        tmp_path / "r.csv", tmp_path / "v.csv", tmp_path / "m.json", tmp_path / "o.xlsx", max_rows=4
    )
    wb = load_workbook(tmp_path / "o.xlsx", read_only=True)
    assert wb.sheetnames == ["roads", "vertices", "vertices_2", "vertices_3", "metadata"]
    assert report["truncated_cells"] == 1 and report["vertices_rows"] == 7
    cell = list(wb["roads"].iter_rows(min_row=2, max_row=2))[0][
        schema.OUTPUT_COLUMNS.index("GEOMETRY_WKT")
    ].value
    assert cell.endswith("…TRUNCATED") and len(cell) <= EXCEL_MAX_CELL_CHARS


def test_na_like_text_survives_the_xlsx_conversion(tmp_path):
    """`keep_default_na=False` is load-bearing, not incidental.

    With pandas' defaults these five literals all become missing values, so a road actually
    named or labelled with one would read correctly in the CSV and be blank in the workbook.
    Only a genuinely empty cell may be blank.
    """
    literals = ["nan", "NA", "null", "None", "N/A"]
    rows = []
    for i, text in enumerate(literals, start=1):
        row = _road_row(i)
        row["ROAD_NAME"] = text
        rows.append(row)
    blank = _road_row(len(literals) + 1)
    blank["ROAD_NAME"] = None
    rows.append(blank)

    write_csv(rows, schema.OUTPUT_COLUMNS, tmp_path / "r.csv")
    write_csv([], schema.VERTEX_COLUMNS, tmp_path / "v.csv")
    write_json({}, tmp_path / "m.json")
    csv_to_xlsx(tmp_path / "r.csv", tmp_path / "v.csv", tmp_path / "m.json", tmp_path / "o.xlsx")

    ws = load_workbook(tmp_path / "o.xlsx", read_only=True)["roads"]
    col = schema.OUTPUT_COLUMNS.index("ROAD_NAME")
    values = [list(r)[col].value for r in ws.iter_rows(min_row=2)]
    assert values[: len(literals)] == literals, "NA-like text must survive as text"
    assert values[len(literals)] is None, "a genuinely empty cell must stay blank"


def test_split_vertices_sheets_reproduce_the_csv_exactly(tmp_path):
    """The union of the vertices sheets must equal the CSV: no row lost, none duplicated.

    `report["vertices_rows"]` counts rows READ, not rows written, so an off-by-one at a sheet
    seam is invisible to it — a mutant overlapping the seam by one row passes the existing test.
    """
    rows = _vertex_rows(7)
    write_csv([_road_row(1)], schema.OUTPUT_COLUMNS, tmp_path / "r.csv")
    write_csv(rows, schema.VERTEX_COLUMNS, tmp_path / "v.csv")
    write_json({}, tmp_path / "m.json")
    csv_to_xlsx(
        tmp_path / "r.csv", tmp_path / "v.csv", tmp_path / "m.json", tmp_path / "o.xlsx", max_rows=4
    )

    wb = load_workbook(tmp_path / "o.xlsx", read_only=True)
    sheets = [s for s in wb.sheetnames if s.startswith("vertices")]
    assert sheets == ["vertices", "vertices_2", "vertices_3"]
    written = []
    for name in sheets:
        ws = wb[name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert header == schema.VERTEX_COLUMNS
        written.extend(tuple(c.value for c in row) for row in ws.iter_rows(min_row=2))
    assert len(written) == len(rows), "every vertex row exactly once across the sheets"
    assert len({(r[0], r[5], r[6]) for r in written}) == len(rows), "no duplicated row at a seam"
    assert [r[6] for r in written] == [row["SEQ"] for row in rows], "order preserved across sheets"


def test_csv_to_xlsx_refuses_swapped_inputs(tmp_path):
    """to-xlsx must never write a workbook whose 'roads' sheet holds vertices."""
    roads = write_csv([_road_row(1)], schema.OUTPUT_COLUMNS, tmp_path / "roads.csv")
    verts = write_csv([dict.fromkeys(schema.VERTEX_COLUMNS, "1")], schema.VERTEX_COLUMNS, tmp_path / "v.csv")
    meta = write_json({}, tmp_path / "metadata.json")
    with pytest.raises(ValueError, match="swapped"):
        csv_to_xlsx(verts, roads, meta, tmp_path / "out.xlsx")
    assert not (tmp_path / "out.xlsx").exists()


def test_csv_to_xlsx_refuses_a_missing_metadata_file(tmp_path):
    roads = write_csv([_road_row(1)], schema.OUTPUT_COLUMNS, tmp_path / "roads.csv")
    verts = write_csv([dict.fromkeys(schema.VERTEX_COLUMNS, "1")], schema.VERTEX_COLUMNS, tmp_path / "v.csv")
    with pytest.raises(ValueError, match="metadata file not found"):
        csv_to_xlsx(roads, verts, tmp_path / "missing.json", tmp_path / "out.xlsx")
    assert not (tmp_path / "out.xlsx").exists()
